"""Inlezen van Excel- en CSV-afschriften.

Banken leveren allemaal een andere kolomindeling. Daarom leest deze module het
bestand eerst uit als een tabel met kopregel, raadt ze zelf welke kolom welk
veld is, en toont ze die keuze ter bevestiging. De bevestigde indeling kan als
profiel bewaard worden zodat een volgende upload van dezelfde bank meteen goed
staat.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

BESTANDSTYPES = {".xlsx", ".xlsm", ".csv", ".txt", ".tsv"}

# De velden waarop een kolom kan worden afgebeeld.
VELDEN = [
    ("boekdatum", "Boekingsdatum", True),
    ("valutadatum", "Valutadatum", False),
    ("verrichtingsdatum", "Verrichtingsdatum", False),
    ("referentie", "Referentie van de bank", False),
    ("beschrijving", "Beschrijving (soort verrichting)", False),
    ("bedrag", "Bedrag", True),
    ("munt", "Munt", False),
    ("eigen_rekening", "Eigen rekeningnummer", False),
    ("tegenpartij_rekening", "Rekening tegenpartij", False),
    ("tegenpartij_naam", "Naam tegenpartij", False),
    ("begunstigde", "Naam begunstigde", False),
    ("mededeling", "Mededeling", False),
    ("mededeling_2", "Tweede mededeling", False),
    ("mededeling_3", "Derde mededeling", False),
]

VERPLICHT = [naam for naam, _, verplicht in VELDEN if verplicht]

# Trefwoorden per veld, gebruikt om kolommen automatisch te herkennen.
TREFWOORDEN = {
    "verrichtingsdatum": ["verrichtingsdatum", "datum verrichting", "transaction date"],
    "referentie": ["referentie", "reference", "referte", "volgnummer"],
    "beschrijving": ["beschrijving", "verrichting", "type verrichting", "omschrijving type",
                     "soort", "transactietype"],
    "boekdatum": ["boekingsdatum", "boekdatum", "datum verrichting", "uitvoeringsdatum",
                  "transactiedatum", "date", "datum"],
    "valutadatum": ["valutadatum", "valuta datum", "value date", "valeur"],
    "bedrag": ["bedrag", "amount", "montant", "transactiebedrag", "bedrag van de verrichting",
               "debet credit", "som"],
    "munt": ["munt", "valuta", "devise", "currency", "muntsoort"],
    "eigen_rekening": ["rekeningnummer", "rekening", "eigen rekening", "iban rekening",
                       "uw rekening", "account", "rekening iban"],
    "tegenpartij_rekening": ["tegenpartij rekening", "rekening tegenpartij", "iban tegenpartij",
                             "tegenrekening", "rekening begunstigde", "counterparty account",
                             "compte contrepartie"],
    "tegenpartij_naam": ["naam tegenpartij", "tegenpartij", "naam begunstigde tegenpartij",
                         "counterparty", "contrepartie", "naam van de tegenpartij"],
    "begunstigde": ["begunstigde", "naam begunstigde", "beneficiary", "beneficiaire",
                    "naam van de begunstigde"],
    "mededeling": ["mededeling", "vrije mededeling", "gestructureerde mededeling",
                   "detail", "communication", "mededeling 1", "vrije mededeling 1"],
    "mededeling_2": ["mededeling 2", "vrije mededeling 2", "detail 2", "omschrijving 2"],
    "mededeling_3": ["mededeling 3", "vrije mededeling 3", "detail 3", "omschrijving 3"],
}


@dataclass
class Kolomprofiel:
    naam: str
    mapping: dict[str, str] = field(default_factory=dict)
    scheidingsteken: str = ";"
    decimaal: str = ","


# Startprofielen voor courante Belgische banken. De kolomnamen zijn een
# vertrekpunt; wijk gerust af in het scherm bij het importeren.
PROFIELEN: dict[str, Kolomprofiel] = {
    "automatisch": Kolomprofiel("Automatisch herkennen"),
    "kbc": Kolomprofiel("KBC", {
        "boekdatum": "Datum", "valutadatum": "Valuta", "bedrag": "Bedrag",
        "munt": "Munt", "eigen_rekening": "Rekeningnummer",
        "tegenpartij_rekening": "Rekening tegenpartij",
        "tegenpartij_naam": "Naam tegenpartij", "mededeling": "Vrije mededeling",
        "mededeling_2": "Gestructureerde mededeling",
    }),
    "belfius": Kolomprofiel("Belfius", {
        "boekdatum": "Boekingsdatum", "valutadatum": "Valutadatum",
        "bedrag": "Bedrag", "munt": "Devies", "eigen_rekening": "Rekening",
        "tegenpartij_rekening": "Rekening tegenpartij",
        "tegenpartij_naam": "Naam tegenpartij", "mededeling": "Mededeling",
    }),
    "ing": Kolomprofiel("ING", {
        "boekdatum": "Boekingsdatum", "valutadatum": "Valutadatum",
        "bedrag": "Bedrag", "munt": "Munteenheid", "eigen_rekening": "Rekeningnummer",
        "tegenpartij_rekening": "Rekening tegenpartij",
        "tegenpartij_naam": "Naam tegenpartij", "mededeling": "Mededeling",
    }),
    "argenta": Kolomprofiel("Argenta", {
        "boekdatum": "Boekdatum",
        "valutadatum": "Valutadatum",
        "verrichtingsdatum": "Verrichtingsdatum",
        "referentie": "Referentie",
        "beschrijving": "Beschrijving",
        "bedrag": "Bedrag",
        "munt": "Munt",
        "eigen_rekening": "Rekening",
        "tegenpartij_rekening": "Rekening tegenpartij",
        "tegenpartij_naam": "Naam tegenpartij",
        "mededeling": "Mededeling",
    }),
}


# --------------------------------------------------------------------------
# Lezen
# --------------------------------------------------------------------------

def _lees_xlsx(pad: Path, blad: str | None = None) -> tuple[list[str], list[list]]:
    from openpyxl import load_workbook

    wb = load_workbook(pad, read_only=True, data_only=True)
    ws = wb[blad] if blad and blad in wb.sheetnames else wb[wb.sheetnames[0]]
    rijen = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rijen:
        return [], []

    # De kopregel is de eerste rij met minstens twee niet-lege tekstcellen.
    kop_index = 0
    for i, rij in enumerate(rijen[:15]):
        gevuld = [c for c in rij if c not in (None, "")]
        if len(gevuld) >= 2 and sum(isinstance(c, str) for c in gevuld) >= 2:
            kop_index = i
            break
    kop = [str(c).strip() if c is not None else f"kolom_{j + 1}"
           for j, c in enumerate(rijen[kop_index])]
    data = [r for r in rijen[kop_index + 1:] if any(c not in (None, "") for c in r)]
    return kop, data


def _lees_csv(pad: Path, scheidingsteken: str | None = None) -> tuple[list[str], list[list]]:
    ruw = pad.read_bytes()
    for codering in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            tekst = ruw.decode(codering)
            break
        except UnicodeDecodeError:
            continue
    else:
        tekst = ruw.decode("latin-1", errors="replace")

    if scheidingsteken is None:
        monster = "\n".join(tekst.splitlines()[:10])
        try:
            scheidingsteken = csv.Sniffer().sniff(monster, delimiters=";,\t|").delimiter
        except csv.Error:
            scheidingsteken = ";" if monster.count(";") >= monster.count(",") else ","

    lezer = csv.reader(io.StringIO(tekst), delimiter=scheidingsteken)
    rijen = [r for r in lezer if any(c.strip() for c in r)]
    if not rijen:
        return [], []
    kop = [c.strip() or f"kolom_{i + 1}" for i, c in enumerate(rijen[0])]
    return kop, rijen[1:]


def lees_bestand(pad: Path, scheidingsteken: str | None = None) -> tuple[list[str], list[list]]:
    achtervoegsel = pad.suffix.lower()
    if achtervoegsel in (".xlsx", ".xlsm"):
        return _lees_xlsx(pad)
    return _lees_csv(pad, scheidingsteken)


# --------------------------------------------------------------------------
# Kolommen herkennen
# --------------------------------------------------------------------------

def _schoon(tekst: str) -> str:
    tekst = re.sub(r"[^a-z0-9 ]", " ", str(tekst).lower())
    return re.sub(r"\s+", " ", tekst).strip()


def detecteer_mapping(kop: list[str], profiel: str = "automatisch") -> dict[str, str]:
    """Wijst elke kolomnaam toe aan een veld. Geeft {veld: kolomnaam}."""
    mapping: dict[str, str] = {}
    schoon_kop = {_schoon(k): k for k in kop}

    voorkeur = PROFIELEN.get(profiel)
    if voorkeur and voorkeur.mapping:
        for veld, kolom in voorkeur.mapping.items():
            gevonden = schoon_kop.get(_schoon(kolom))
            if gevonden:
                mapping[veld] = gevonden

    gebruikt = set(mapping.values())
    for veld, trefwoorden in TREFWOORDEN.items():
        if veld in mapping:
            continue
        beste = None
        beste_score = 0
        for schoon, origineel in schoon_kop.items():
            if origineel in gebruikt:
                continue
            for trefwoord in trefwoorden:
                tw = _schoon(trefwoord)
                if schoon == tw:
                    score = 100
                elif tw in schoon:
                    score = 70 + len(tw)
                elif schoon in tw:
                    score = 40 + len(schoon)
                else:
                    continue
                if score > beste_score:
                    beste, beste_score = origineel, score
        if beste:
            mapping[veld] = beste
            gebruikt.add(beste)

    # Sommige banken splitsen debet en credit over twee kolommen.
    if "bedrag" not in mapping:
        debet = next((k for k in kop if _schoon(k) in ("debet", "uitgaven", "af")), None)
        credit = next((k for k in kop if _schoon(k) in ("credit", "inkomsten", "bij")), None)
        if debet and credit:
            mapping["bedrag_debet"] = debet
            mapping["bedrag_credit"] = credit
    return mapping


# --------------------------------------------------------------------------
# Waarden omzetten
# --------------------------------------------------------------------------

_DATUMFORMATEN = [
    "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y/%m/%d",
    "%d/%m/%y", "%d-%m-%y", "%m/%d/%Y", "%d %b %Y", "%Y%m%d",
]


def parse_datum(waarde) -> date | None:
    if waarde in (None, ""):
        return None
    if isinstance(waarde, datetime):
        return waarde.date()
    if isinstance(waarde, date):
        return waarde
    tekst = str(waarde).strip()
    tekst = re.sub(r"\s+\d{1,2}:\d{2}(:\d{2})?$", "", tekst)
    for opmaak in _DATUMFORMATEN:
        try:
            return datetime.strptime(tekst, opmaak).date()
        except ValueError:
            continue
    return None


def parse_bedrag(waarde, decimaal: str = "auto") -> Decimal | None:
    if waarde in (None, ""):
        return None
    if isinstance(waarde, (int, float, Decimal)):
        return Decimal(str(waarde))

    tekst = str(waarde).strip()
    negatief = tekst.startswith("(") and tekst.endswith(")")
    tekst = tekst.strip("()")
    tekst = re.sub(r"[^\d,.\-+]", "", tekst)
    if not tekst:
        return None

    if decimaal == "auto":
        laatste_komma = tekst.rfind(",")
        laatste_punt = tekst.rfind(".")
        if laatste_komma > laatste_punt:
            tekst = tekst.replace(".", "").replace(",", ".")
        else:
            tekst = tekst.replace(",", "")
    elif decimaal == ",":
        tekst = tekst.replace(".", "").replace(",", ".")
    else:
        tekst = tekst.replace(",", "")

    try:
        bedrag = Decimal(tekst)
    except InvalidOperation:
        return None
    return -bedrag if negatief else bedrag


def rij_naar_velden(rij: list, kop: list[str], mapping: dict[str, str],
                    decimaal: str = "auto") -> dict:
    """Zet één tabelrij om naar een woordenboek met genormaliseerde velden."""
    index = {naam: i for i, naam in enumerate(kop)}

    def haal(veld: str):
        kolom = mapping.get(veld)
        if not kolom or kolom not in index:
            return None
        i = index[kolom]
        return rij[i] if i < len(rij) else None

    bedrag = parse_bedrag(haal("bedrag"), decimaal)
    if bedrag is None and "bedrag_debet" in mapping:
        debet = parse_bedrag(haal("bedrag_debet"), decimaal) or Decimal("0")
        credit = parse_bedrag(haal("bedrag_credit"), decimaal) or Decimal("0")
        bedrag = credit - abs(debet)

    mededelingen = [
        str(haal(v)).strip()
        for v in ("mededeling", "mededeling_2", "mededeling_3")
        if haal(v) not in (None, "")
    ]

    def tekst(veld):
        w = haal(veld)
        return str(w).strip() if w not in (None, "") else ""

    return {
        "boekdatum": parse_datum(haal("boekdatum")),
        "valutadatum": parse_datum(haal("valutadatum")),
        "verrichtingsdatum": parse_datum(haal("verrichtingsdatum")),
        "referentie": tekst("referentie"),
        "beschrijving": tekst("beschrijving"),
        "bedrag": bedrag,
        "munt": tekst("munt") or "EUR",
        "eigen_rekening": tekst("eigen_rekening"),
        "tegenpartij_rekening": tekst("tegenpartij_rekening"),
        "tegenpartij_naam": tekst("tegenpartij_naam"),
        "begunstigde": tekst("begunstigde"),
        "mededeling": " | ".join(mededelingen),
        "ruw": {k: (str(v) if v is not None else "") for k, v in zip(kop, rij)},
    }
